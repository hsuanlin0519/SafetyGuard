import os
import json
import csv
import docx
import pandas as pd
import PyPDF2
import xml.etree.ElementTree as ET
import openpyxl
import pdfplumber
import olefile
import xlrd
import win32com.client as win32


def extract_text(file_path):
    try:
        # 取得副檔名
        _, file_extension = os.path.splitext(file_path)
        file_extension = file_extension.lower()

        text_content = ""
        bool_image = False
        if file_extension == '.json':
            # 處理 JSON 檔案
            with open(file_path, 'r', encoding='utf-8') as f:
                data = json.load(f)
                text_content = json.dumps(data, ensure_ascii=False, indent=2)
        elif file_extension == '.jsonl':
            # 處理 JSONL 檔案
            with open(file_path, 'r', encoding='utf-8') as f:
                lines = f.readlines()
                text_content = "\n".join(lines)

        elif file_extension == '.xml':
            # 處理 XML 檔案
            tree = ET.parse(file_path)
            root = tree.getroot()
            text_content = ET.tostring(root, encoding='unicode')

        elif file_extension == '.docx':
            # 處理 DOCX 檔案
            doc = docx.Document(file_path)
            text_content = "\n".join([para.text for para in doc.paragraphs])
            for rel in doc.part.rels.values():
                if "image" in rel.target_ref:
                    bool_image = True

        elif file_extension == '.doc':
            # 處理 DOC 檔案
            word = win32.Dispatch("Word.Application")
            doc = word.Documents.Open(file_path)
            # 提取文字
            text_content = doc.Content.Text
            # 檢測圖片
            bool_image = False
            for shape in doc.InlineShapes:
                if shape.Type == 3:  # 3 是圖片類型
                    bool_image = True
                    break
            doc.Close()
            word.Quit()

        elif file_extension == '.xlsx':
            # 處理 Excel 檔案
            wb = openpyxl.load_workbook(file_path)
            text_lines = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    # 將每行中的非空單元格轉換為字串並連接
                    text_line = ' '.join(str(cell) for cell in row if cell is not None)
                    text_lines.append(text_line)
                if sheet._images:
                    bool_image = True  # 找到圖片
                    break

            # 將所有行連接為一個大字串
            text_content = '\n'.join(text_lines)

        elif file_extension == '.xls':
            wb = xlrd.open_workbook(file_path)
            text_lines = []
            # 遍歷所有工作表
            for sheet in wb.sheets():
                # 遍歷工作表中的每一行和每一列
                for row_idx in range(sheet.nrows):
                    row_values = [str(sheet.cell_value(row_idx, col_idx)) for col_idx in range(sheet.ncols)]
                    # 將每行中的文字合併並加入 text_lines
                    text_lines.append(" ".join(row_values))
            # 將所有行連接為一個大字串
            text_content = "\n".join(text_lines)
            # 嘗試打開 OLE 格式的 .xls 文件
            if olefile.isOleFile(file_path):
                ole = olefile.OleFileIO(file_path)
                for entry in ole.listdir():
                    # 檢查是否存在與圖片相關的流
                    if any(substring in entry[0] for substring in ['image', 'Pictures', 'PICTURE']):
                        bool_image = True  # 找到圖片
                        break
                ole.close()
        elif file_extension == '.csv':
            # 處理 CSV 檔案
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                reader = csv.reader(f)
                text_content = "\n".join([", ".join(row) for row in reader])

        elif file_extension == '.txt':
            # 處理 TXT 檔案
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                text_content = f.read()

        elif file_extension == '.pdf':
            # 處理 PDF 檔案
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages:
                    text_content += page.extract_text()
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    # 檢查頁面中是否包含圖片
                    if page.images:
                        bool_image = True
                        break
        else:
            return True, None, f"Unsupported file type: {file_extension}"

        return True, bool_image, text_content

    except Exception as e:
        return False, None, f"Error processing file: {str(e)}"
